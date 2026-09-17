import openpyxl

files = [
    "../college_student_queries.xlsx",
    "../college_student_queries_v2.xlsx"
]

for f_path in files:
    try:
        wb = openpyxl.load_workbook(f_path, data_only=True)
        print(f"\n==========================================")
        print(f"Dumping file: {f_path}")
        print(f"Sheets: {wb.sheetnames}")
        print(f"==========================================")
        for name in wb.sheetnames:
            sheet = wb[name]
            print(f"--- Sheet: {name} ({sheet.max_row} rows) ---")
            for r_idx in range(1, min(100, sheet.max_row + 1)):
                row = [sheet.cell(row=r_idx, column=c_idx).value for c_idx in range(1, sheet.max_column + 1)]
                if any(cell is not None for cell in row):
                    row_str = " | ".join([str(c) for c in row if c is not None])
                    print(f"Row {r_idx}: {row_str[:200]}")
    except Exception as e:
        print(f"Error reading {f_path}: {e}")
