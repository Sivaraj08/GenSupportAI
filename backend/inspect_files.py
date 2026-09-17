import os
import openpyxl
from pypdf import PdfReader

files = [
    "../college_student_queries.xlsx",
    "../college_student_queries_v2.xlsx",
    "../college_student_queries.pdf"
]

for f_path in files:
    abs_path = os.path.abspath(f_path)
    if not os.path.exists(abs_path):
        print(f"File not found: {f_path}")
        continue
    
    print(f"\n==========================================")
    print(f"Inspecting file: {f_path}")
    print(f"==========================================")
    
    ext = os.path.splitext(f_path)[1].lower()
    if ext == ".xlsx":
        try:
            wb = openpyxl.load_workbook(abs_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                print(f"Sheet: {sheet_name}")
                sheet = wb[sheet_name]
                for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                    row_str = " | ".join([str(cell) for cell in row if cell is not None])
                    if any(word in row_str.lower() for word in ["mark", "assess", "weight"]):
                        print(f"  Row {r_idx+1}: {row_str[:300]}")
        except Exception as e:
            print(f"Error reading excel: {e}")
            
    elif ext == ".pdf":
        try:
            reader = PdfReader(abs_path)
            for p_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if not text:
                    continue
                lines = text.split("\n")
                for l_idx, line in enumerate(lines):
                    if any(word in line.lower() for word in ["mark", "assess", "weight"]):
                        print(f"  Page {p_idx+1}, Line {l_idx+1}: {line}")
        except Exception as e:
            print(f"Error reading pdf: {e}")
