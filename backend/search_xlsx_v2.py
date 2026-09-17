import openpyxl

wb = openpyxl.load_workbook("../college_student_queries_v2.xlsx", data_only=True)
print("college_student_queries_v2.xlsx sheets:", wb.sheetnames)
for name in wb.sheetnames:
    sheet = wb[name]
    print(f"\nSearching sheet: {name}")
    for r_idx in range(1, sheet.max_row + 1):
        row = [sheet.cell(row=r_idx, column=c_idx).value for c_idx in range(1, sheet.max_column + 1)]
        row_str = " | ".join([str(c) for c in row if c is not None])
        if any(word in row_str.lower() for word in ["mark", "assess", "weight"]):
            print(f"Row {r_idx}: {row_str}")
